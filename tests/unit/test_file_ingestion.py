"""Tests for file (pptx/docx/pdf) ingestion into DocumentIR."""

from pathlib import Path

import pytest

from animate_agent.documents.file_parser import (
    parse_docx,
    parse_file,
    parse_markdown,
    parse_pdf,
    parse_pptx,
)
from animate_agent.documents.models import DocumentIR


def _write_pptx(path: Path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "力学基础"
    presentation.save(str(path))


def _write_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("电路基础", level=1)
    document.add_paragraph("电压与电流的关系")
    document.add_heading("欧姆定律", level=1)
    document.add_paragraph("U = I * R")
    document.save(str(path))


def _write_pdf(path: Path) -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as handle:
        writer.write(handle)


def test_parse_pptx_extracts_text(tmp_path: Path) -> None:
    path = tmp_path / "sample.pptx"
    _write_pptx(path)

    document = parse_pptx(path)

    assert document.source.type == "file"
    assert document.document_id
    assert document.sections[0].title == "力学基础"


def test_parse_docx_groups_headings(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    _write_docx(path)

    document = parse_docx(path)

    assert document.source.type == "file"
    assert [section.title for section in document.sections if section.title] == [
        "电路基础",
        "欧姆定律",
    ]
    assert any(
        block.type == "paragraph" and "电压" in block.text
        for section in document.sections
        for block in section.blocks
    )


def test_parse_pdf_produces_document(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    _write_pdf(path)

    document = parse_pdf(path)

    assert document.source.type == "file"
    assert document.document_id


def test_parse_file_dispatches(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    _write_docx(path)

    assert isinstance(parse_file(path), DocumentIR)


def test_parse_file_rejects_unknown_extension(tmp_path: Path) -> None:
    path = tmp_path / "sample.xyz"
    path.write_text("hello", encoding="utf-8")

    with pytest.raises(ValueError):
        parse_file(path)


MARKDOWN_SAMPLE = """---
title: 前端元数据
---

# 力学基础

物体在不受外力时保持静止。

## 工具准备

- 木工锯
- 木工凿
- 木工刨

```python
print("hello")
```
"""


def test_parse_markdown_splits_headings_into_sections(tmp_path: Path) -> None:
    path = tmp_path / "sample.md"
    path.write_text(MARKDOWN_SAMPLE, encoding="utf-8")

    document = parse_markdown(path)

    assert document.source.type == "file"
    assert [section.title for section in document.sections] == ["力学基础", "工具准备"]
    assert [section.level for section in document.sections] == [1, 2]


def test_parse_markdown_skips_frontmatter(tmp_path: Path) -> None:
    path = tmp_path / "sample.md"
    path.write_text(MARKDOWN_SAMPLE, encoding="utf-8")

    text = " ".join(
        block.text for section in parse_markdown(path).sections for block in section.blocks
    )

    assert "前端元数据" not in text


def test_parse_markdown_builds_paragraph_list_and_code_blocks(tmp_path: Path) -> None:
    path = tmp_path / "sample.md"
    path.write_text(MARKDOWN_SAMPLE, encoding="utf-8")

    kinds = {block.type for section in parse_markdown(path).sections for block in section.blocks}

    assert kinds == {"paragraph", "list", "code"}


def test_parse_markdown_strips_inline_markers(tmp_path: Path) -> None:
    path = tmp_path / "sample.md"
    path.write_text(
        "# 标题\n\n这是 **加粗** 和 [链接](https://example.com) 文本。\n",
        encoding="utf-8",
    )

    text = parse_markdown(path).sections[0].blocks[0].text

    assert text == "这是 加粗 和 链接 文本。"


def test_parse_markdown_handles_plain_text_without_headings(tmp_path: Path) -> None:
    path = tmp_path / "sample.txt"
    path.write_text("第一段内容。\n\n第二段内容。\n", encoding="utf-8")

    document = parse_markdown(path)

    assert len(document.sections) == 1
    assert [block.text for block in document.sections[0].blocks] == ["第一段内容。", "第二段内容。"]


def test_parse_markdown_rejects_empty_document(tmp_path: Path) -> None:
    path = tmp_path / "empty.md"
    path.write_text("\n\n   \n", encoding="utf-8")

    with pytest.raises(ValueError, match="没有可解析"):
        parse_markdown(path)


PLAIN_TEXT_SAMPLE = """抛体运动入门教程

【一、基础知识介绍】

1. 定义：把物体以一定的初速度抛出。

【二、解题步骤】

- 建立坐标系。
"""


def test_parse_markdown_recognizes_bracket_headings(tmp_path: Path) -> None:
    path = tmp_path / "sample.txt"
    path.write_text(PLAIN_TEXT_SAMPLE, encoding="utf-8")

    document = parse_markdown(path)

    assert [section.title for section in document.sections if section.title] == [
        "一、基础知识介绍",
        "二、解题步骤",
    ]


def test_parse_markdown_uses_leading_line_as_title(tmp_path: Path) -> None:
    path = tmp_path / "sample.txt"
    path.write_text(PLAIN_TEXT_SAMPLE, encoding="utf-8")

    document = parse_markdown(path)

    assert document.title == "抛体运动入门教程"


def test_parse_markdown_drops_promoted_title_line_from_blocks(tmp_path: Path) -> None:
    path = tmp_path / "sample.txt"
    path.write_text(PLAIN_TEXT_SAMPLE, encoding="utf-8")

    document = parse_markdown(path)

    texts = [block.text for section in document.sections for block in section.blocks]
    assert "抛体运动入门教程" not in texts
    assert all(section.id != "section-overview" for section in document.sections)


def test_parse_markdown_keeps_leading_paragraph_when_no_headings(tmp_path: Path) -> None:
    path = tmp_path / "sample.txt"
    path.write_text("第一段内容。\n\n第二段内容。\n", encoding="utf-8")

    document = parse_markdown(path)

    # Without structure to spare, the leading paragraph is content, not a title.
    assert len(document.sections) == 1
    assert [block.text for block in document.sections[0].blocks] == ["第一段内容。", "第二段内容。"]


def test_parse_markdown_recognizes_chinese_numbered_headings(tmp_path: Path) -> None:
    path = tmp_path / "sample.txt"
    path.write_text("教程\n\n一、基础\n\n正文一。\n\n二、进阶\n\n正文二。\n", encoding="utf-8")

    document = parse_markdown(path)

    assert [section.title for section in document.sections if section.title] == [
        "一、基础",
        "二、进阶",
    ]


def test_parse_markdown_does_not_treat_long_paragraph_as_heading(tmp_path: Path) -> None:
    path = tmp_path / "sample.txt"
    path.write_text(
        "标题\n\n【一、这是一段很长的正文，长到超过了标题长度上限，"
        "因此不该被当成章节标题，而应原样保留为正文段落。】\n",
        encoding="utf-8",
    )

    document = parse_markdown(path)

    assert [section.title for section in document.sections if section.title] == []
    texts = [block.text for section in document.sections for block in section.blocks]
    assert any("很长的正文" in text for text in texts)


def test_parse_file_dispatches_markdown_and_text(tmp_path: Path) -> None:
    md = tmp_path / "sample.md"
    md.write_text("# 标题\n\n正文。\n", encoding="utf-8")
    txt = tmp_path / "sample.txt"
    txt.write_text("正文。\n", encoding="utf-8")

    assert isinstance(parse_file(md), DocumentIR)
    assert isinstance(parse_file(txt), DocumentIR)
