"""
文件文本提取器 — 仅提取纯文本，不保留格式。
支持 .pptx / .docx / .pdf
"""

import os
from pptx import Presentation
from docx import Document
import pdfplumber


def read_pptx(path: str) -> str:
    """
    读取 .pptx 文件，返回所有幻灯片中的纯文本。
    每页幻灯片文本之间用空行分隔。
    """
    prs = Presentation(path)
    slides_text: list[str] = []

    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        lines.append(text)
        if lines:
            slides_text.append(f"--- 第 {i} 页 ---\n" + "\n".join(lines))

    return "\n\n".join(slides_text)


def read_docx(path: str) -> str:
    """
    读取 .docx 文件，返回所有段落的纯文本。
    空段落会被跳过。
    """
    doc = Document(path)
    paragraphs: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    return "\n".join(paragraphs)


def read_pdf(path: str) -> str:
    """
    读取 .pdf 文件，返回所有页面的纯文本。
    每页文本之间用空行分隔。
    """
    with pdfplumber.open(path) as pdf:
        pages_text: list[str] = []
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text()
            if text:
                pages_text.append(f"--- 第 {i} 页 ---\n{text.strip()}")
        return "\n\n".join(pages_text)


def extract_text(path: str) -> str:
    """
    自动识别文件格式并提取纯文本。
    支持 .pptx / .docx / .pdf
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pptx":
        return read_pptx(path)
    elif ext == ".docx":
        return read_docx(path)
    elif ext == ".pdf":
        return read_pdf(path)
    else:
        raise ValueError(f"不支持的文件格式: {ext}。支持: .pptx, .docx, .pdf")
