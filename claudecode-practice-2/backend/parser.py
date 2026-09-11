"""
文件解析器 — 提取结构化文档（DocumentIR）与纯文本。
支持 .pptx / .docx / .pdf

- read_pptx / read_docx / read_pdf(path) -> DocumentIR   （结构化）
- extract_document_ir(path) -> DocumentIR                 （自动识别格式）
- extract_text(path) -> str                               （纯文本，兼容旧接口）
"""

import os
from pptx import Presentation
from docx import Document
import pdfplumber

from ir_schemas import DocumentIR, Section


def _is_heading_style(style_name: str) -> bool:
    """判断 Word 段落样式是否为标题样式（英文 Heading N / 中文 标题 N / Title）。"""
    s = (style_name or "").lower()
    return s.startswith("heading") or "标题" in style_name or s == "title"


def _resolve_title(meta_title: str, sections: list) -> str:
    """标题回退链：元数据标题 → 首节标题 → 首节第一段。"""
    if meta_title:
        return meta_title
    if not sections:
        return ""
    return sections[0].heading or (sections[0].paragraphs[0] if sections[0].paragraphs else "")


def read_pptx(path: str) -> DocumentIR:
    """读取 .pptx，返回结构化 DocumentIR。"""
    prs = Presentation(path)
    title = (prs.core_properties.title or "").strip()
    sections: list[Section] = []
    slides_text: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        lines.append(text)
        if lines:
            slides_text.append(f"--- 第 {i} 页 ---\n" + "\n".join(lines))
            sections.append(Section(heading=lines[0], paragraphs=lines[1:]))
    raw_text = "\n\n".join(slides_text)
    return DocumentIR(
        source_type="pptx",
        title=_resolve_title(title, sections),
        sections=sections,
        raw_text=raw_text,
    )


def read_docx(path: str) -> DocumentIR:
    """读取 .docx，返回结构化 DocumentIR。"""
    doc = Document(path)
    title = (doc.core_properties.title or "").strip()
    sections: list[Section] = []
    paragraphs: list[str] = []
    cur_heading = ""
    cur_paras: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        paragraphs.append(text)
        style_name = para.style.name if para.style is not None else ""
        if _is_heading_style(style_name):
            if cur_heading or cur_paras:
                sections.append(Section(heading=cur_heading, paragraphs=cur_paras))
            cur_heading = text
            cur_paras = []
        else:
            cur_paras.append(text)
    if cur_heading or cur_paras:
        sections.append(Section(heading=cur_heading, paragraphs=cur_paras))
    raw_text = "\n".join(paragraphs)
    return DocumentIR(
        source_type="docx",
        title=_resolve_title(title, sections),
        sections=sections,
        raw_text=raw_text,
    )


def read_pdf(path: str) -> DocumentIR:
    """读取 .pdf，返回结构化 DocumentIR。"""
    with pdfplumber.open(path) as pdf:
        meta_title = ((pdf.metadata or {}).get("Title") or "").strip()
        sections: list[Section] = []
        pages_text: list[str] = []
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text()
            if not text:
                continue
            pages_text.append(f"--- 第 {i} 页 ---\n{text.strip()}")
            lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
            sections.append(Section(heading=lines[0] if lines else "", paragraphs=lines[1:]))
        raw_text = "\n\n".join(pages_text)
    return DocumentIR(
        source_type="pdf",
        title=_resolve_title(meta_title, sections),
        sections=sections,
        raw_text=raw_text,
    )


def extract_document_ir(path: str) -> DocumentIR:
    """自动识别文件格式，提取为结构化 DocumentIR。支持 .pptx / .docx / .pdf"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pptx":
        return read_pptx(path)
    elif ext == ".docx":
        return read_docx(path)
    elif ext == ".pdf":
        return read_pdf(path)
    else:
        raise ValueError(f"不支持的文件格式: {ext}。支持: .pptx, .docx, .pdf")


def extract_text(path: str) -> str:
    """自动识别文件格式并提取纯文本。支持 .pptx / .docx / .pdf（兼容旧接口）。"""
    return extract_document_ir(path).raw_text
